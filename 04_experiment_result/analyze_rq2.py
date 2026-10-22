#!/usr/bin/env python3
"""Reproduce RQ2 statistics from bundled evaluation records, without model calls."""
import argparse
import csv
import json
import statistics
from collections import Counter, defaultdict
from fractions import Fraction
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ARMS = ['Enhanced UML', 'Raw NL', 'Spec-driven NL']
MODELS = ['claude-opus-4-6', 'claude-sonnet-4-6']
MAX_LLM_CALLS = 50


def handlers():
    return json.loads((HERE / 'RQ2_handlers.json').read_text(encoding='utf-8'))


def load_runs():
    """Read evaluation records and compute run-level CPR."""
    expected = {(h['layer'], h['eval_id']): h['expected_ut_total'] for h in handlers()}
    book = load_workbook(HERE / 'RQ2_design_artifacts.xlsx', read_only=True, data_only=True)
    result = []
    for arm in ARMS:
        iterator = book[arm].iter_rows(values_only=True)
        headers = next(iterator)
        for values in iterator:
            row = dict(zip(headers, values))
            assert row['arm'] == arm
            total = expected[row['layer'], row['eval_id']]
            assert row['expected_ut_total'] == total and 0 <= row['ut_passed'] <= total
            assert row['compile_success'] or row['ut_passed'] == 0
            assert 0 <= row['llm_calls'] <= MAX_LLM_CALLS and row['token_total'] >= 0
            row['cpr'] = row['ut_passed'] / total
            result.append(row)
    book.close()
    assert len(result) == 743
    keys = [(r['arm'], r['model'], r['layer'], r['eval_id'], r['developer'], r['round']) for r in result]
    assert len(keys) == len(set(keys)), 'Duplicate run identity'
    for model in MODELS:
        for arm in ARMS:
            assert {(r['layer'], r['eval_id']) for r in result if (r['model'], r['arm']) == (model, arm)} == set(expected)
    return result


def exact_wilcoxon(differences):
    """One-sided signed-rank test (greater), exact sign permutations, midranks.

    Fractions avoid floating-point zero/tie artifacts. Drop zero differences;
    enumerate the null distribution by dynamic programming over doubled ranks.
    """
    values = [d for d in differences if d]
    if not values:
        return {'statistic': 0.0, 'p': 1.0, 'n_nonzero': 0}
    order = sorted(range(len(values)), key=lambda i: abs(values[i]))
    ranks = [0] * len(values)
    i = 0
    while i < len(order):
        j = i + 1
        while j < len(order) and abs(values[order[j]]) == abs(values[order[i]]):
            j += 1
        for k in order[i:j]:
            ranks[k] = i + 1 + j
        i = j
    observed = sum(rank for rank, value in zip(ranks, values) if value > 0)
    counts = Counter({0: 1})
    for rank in ranks:
        next_counts = counts.copy()
        for total, count in counts.items():
            next_counts[total + rank] += count
        counts = next_counts
    return {'statistic': observed / 2, 'p': sum(n for s, n in counts.items() if s >= observed) / 2 ** len(values),
            'n_nonzero': len(values)}


def holm_adjust(p_values):
    """Holm step-down adjustment for one family, in original input order."""
    adjusted = [0.0] * len(p_values)
    running_max = 0.0
    for rank, index in enumerate(sorted(range(len(p_values)), key=p_values.__getitem__)):
        running_max = max(running_max, (len(p_values) - rank) * p_values[index])
        adjusted[index] = min(1.0, running_max)
    return adjusted


def aggregate(rows):
    groups = defaultdict(list)
    for r in rows:
        groups[(r['model'], r['arm'], r['layer'], r['eval_id'])].append(r)
    by_handler = []
    paired = {}
    for (model, arm, layer, name), group in sorted(groups.items()):
        cpr = statistics.mean(Fraction(r['ut_passed'], r['expected_ut_total']) for r in group)
        tokens = statistics.mean(Fraction(r['token_total']) for r in group)
        paired[model, arm, layer, name] = (cpr, tokens)
        by_handler.append(dict(model=model, arm=arm, layer=layer, eval_id=name, n_runs=len(group),
                               cpr=float(cpr), tokens_k=float(tokens / 1000),
                               llm_calls=statistics.mean(r['llm_calls'] for r in group)))
    summary = []
    for model in MODELS:
        baseline_tokens = statistics.mean(v[1] for k, v in paired.items() if k[:2] == (model, ARMS[0]))
        for arm in ARMS:
            group = [r for r in rows if (r['model'], r['arm']) == (model, arm)]
            cells = [r for r in by_handler if (r['model'], r['arm']) == (model, arm)]
            token_mean = statistics.mean(paired[model, arm, h['layer'], h['eval_id']][1] for h in handlers())
            entry = dict(model=model, arm=arm, n_runs=len(group), n_handlers=len(cells),
                         micro_cpr=sum(r['ut_passed'] for r in group) / sum(r['expected_ut_total'] for r in group),
                         macro_cpr=statistics.mean(r['cpr'] for r in cells), tokens_k=float(token_mean / 1000),
                         token_ratio=float(token_mean / baseline_tokens),
                         llm_calls=statistics.mean(r['llm_calls'] for r in cells),
                         cpr_p=None, tokens_p=None, cpr_p_holm=None, tokens_p_holm=None,
                         cpr_n_nonzero=None, tokens_n_nonzero=None)
            if arm != ARMS[0]:
                for metric, index, direction in [('cpr', 0, -1), ('tokens', 1, 1)]:
                    differences = [direction * (paired[model, arm, h['layer'], h['eval_id']][index] -
                                                paired[model, ARMS[0], h['layer'], h['eval_id']][index]) for h in handlers()]
                    test = exact_wilcoxon(differences)
                    entry[metric + '_p'] = test['p']
                    entry[metric + '_n_nonzero'] = test['n_nonzero']
            summary.append(entry)
    comparisons = [r for r in summary if r['arm'] != ARMS[0]]
    for metric in ('cpr', 'tokens'):
        adjusted = holm_adjust([r[metric + '_p'] for r in comparisons])
        for entry, p in zip(comparisons, adjusted):
            entry[metric + '_p_holm'] = p
    return summary, by_handler


def sample_counts(rows):
    actual = Counter((r['arm'], r['model'], r['layer'], r['eval_id'], r['developer']) for r in rows)
    output = []
    for arm in ARMS:
        for model in MODELS:
            for h in handlers():
                for developer in ([None] if arm == ARMS[0] else ['developer_1', 'developer_2']):
                    n = actual[arm, model, h['layer'], h['eval_id'], developer]
                    output.append(dict(arm=arm, model=model, layer=h['layer'], eval_id=h['eval_id'],
                                       developer=developer, n_runs=n, nominal_runs=5, missing_runs=5-n))
    assert sum(r['missing_runs'] for r in output) == 7
    return output


def save_json(out, name, data):
    (out / name).write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def save_csv(out, name, rows):
    with (out / name).open('w', encoding='utf-8', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def table_rows(summary):
    lines = []
    for r in summary:
        arm = r'\enhancedUML{}' if r['arm'] == ARMS[0] else r['arm']
        p = '--' if r['cpr_p_holm'] is None else f"{r['cpr_p_holm']:.6f} / {r['tokens_p_holm']:.6f}"
        lines.append(f"& {arm} & {r['micro_cpr']:.3f} & {r['macro_cpr']:.3f} & {r['tokens_k']:.1f} & "
                     f"{r['token_ratio']:.2f}" + r'$\times$' + f" & {p} " + r'\\')
    return '\n'.join(lines) + '\n'


def check_paper(path, summary):
    text = path.read_text(encoding='utf-8')
    # Ignore bold formatting and whitespace, preserving every reported value.
    import re
    text = re.sub(r'\\textbf\{([^{}]*)\}', r'\1', text)
    normalize = lambda s: re.sub(r'\s+', '', s)
    text = normalize(text)
    expected = table_rows(summary).splitlines()
    missing = [line for line in expected if normalize(line) not in text]
    if missing:
        raise ValueError('Paper table mismatch:\n' + '\n'.join(missing))
    print('Paper check: all RQ2 execution-table values match.')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir', type=Path, default=HERE / 'analysis/rq2')
    parser.add_argument('--paper-source', type=Path, help='Optionally verify the RQ2 execution table')
    args = parser.parse_args()
    out = args.output_dir
    out.mkdir(parents=True, exist_ok=True)
    rows = load_runs()
    summary, by_handler = aggregate(rows)
    counts = sample_counts(rows)
    for name, values in [('summary', summary), ('by_handler', by_handler), ('sample_counts', counts)]:
        save_json(out, name + '.json', values)
        save_csv(out, name + '.csv', values)
    save_json(out, 'runs.json', rows)
    save_json(out, 'validation.json', {'n_runs': len(rows), 'missing_runs': [r for r in counts if r['missing_runs']],
                                     'agent_call_budget': MAX_LLM_CALLS,
                                     'missing_run_policy': 'analyze available records; do not impute missing runs',
                                     'wilcoxon': 'paired handler means; one-sided exact sign permutations; zero differences dropped, tied absolute differences midranked',
                                     'multiple_comparisons': 'Holm correction separately for CPR and tokens; each family contains four comparisons (two models times two NL baselines)',
                                     'p_value_fields': '*_p: raw; *_p_holm: Holm-adjusted, reported in the paper table'})
    (out / 'table_rows.tex').write_text(table_rows(summary), encoding='utf-8')
    print('RQ2: 743 evaluation records; 7 missing records.')
    for r in summary:
        print(r['model'], r['arm'], f"N={r['n_runs']}, Micro/Macro CPR={r['micro_cpr']:.3f}/{r['macro_cpr']:.3f}, tokens={r['tokens_k']:.1f}K, CPR p(raw/Holm)={r['cpr_p']}/{r['cpr_p_holm']}")
    if args.paper_source:
        check_paper(args.paper_source, summary)


if __name__ == '__main__':
    main()
